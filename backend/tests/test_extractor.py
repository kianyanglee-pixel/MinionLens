from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document

from app.extractor import extract_fields, _text_from_docx, _text_from_pdf


def _fake_inbox(read_text=None, read_bytes=None):
    return SimpleNamespace(
        read_text=lambda path: read_text,
        read_bytes=lambda path: read_bytes,
    )


def test_text_from_docx_extracts_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Shipper: Acme Freight Ltd")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Port of Loading"
    table.rows[0].cells[1].text = "Singapore"
    buf = BytesIO()
    doc.save(buf)

    text = _text_from_docx(buf.getvalue())

    assert "Shipper: Acme Freight Ltd" in text
    assert "Port of Loading" in text
    assert "Singapore" in text


def test_text_from_pdf_joins_page_text_and_skips_none():
    fake_page_1 = MagicMock(extract_text=lambda: "Page one text")
    fake_page_2 = MagicMock(extract_text=lambda: None)
    fake_pdf = MagicMock()
    fake_pdf.pages = [fake_page_1, fake_page_2]
    fake_pdf.__enter__.return_value = fake_pdf

    with patch("app.extractor.pdfplumber.open", return_value=fake_pdf):
        text = _text_from_pdf(b"irrelevant")

    assert text == "Page one text"


def test_extract_fields_rejects_unsupported_format():
    inbox = _fake_inbox()

    result = extract_fields(inbox, "attachments/email_001_SI.pptx")

    assert result["ok"] is False
    assert result["error"] == "rejected_format:.pptx"


def test_extract_fields_pdf_with_no_text_layer_is_unreadable():
    inbox = _fake_inbox(read_bytes=b"scanned-image-only")

    with patch("app.extractor._text_from_pdf", return_value=""):
        result = extract_fields(inbox, "attachments/email_002_BL.pdf")

    assert result == {
        "path": "attachments/email_002_BL.pdf",
        "ok": False,
        "error": "unreadable",
        "fields": {},
        "found": {},
        "text": "",
    }


def test_extract_fields_txt_happy_path_returns_source_text():
    source_text = "Shipper: Acme\nPort of Loading: Singapore"
    inbox = _fake_inbox(read_text=source_text)
    fake_fields = {
        "shipper": "Acme",
        "consignee": None,
        "notify_party": None,
        "port_of_loading": "Singapore",
        "port_of_discharge": None,
        "container_count": None,
        "gross_weight_kg": None,
    }

    with patch("app.extractor.ask_json", return_value=fake_fields):
        result = extract_fields(inbox, "attachments/email_003_SI.txt")

    assert result["ok"] is True
    assert result["fields"]["shipper"] == "Acme"
    assert result["found"]["shipper"] is True
    assert result["found"]["consignee"] is False
    assert result["text"] == source_text
